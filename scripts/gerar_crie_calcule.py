"""Gera assets/crie_calcule.json (materiais do "Crie e Avalie") a partir da planilha do BE³ Build.

- Lista de materiais e impactos (por 1 m² de parede): aba "Impactos do crie e calcule".
- Propriedades térmicas das camadas: aba "Prop. térmicas crie e calcule"
  (mesma leitura do extrair_prop_termicas.py usado com o BE3_Prop_Termicas_CrieCalcule.js).

Uso: python3 scripts/gerar_crie_calcule.py "caminho/app planilha.xlsx"
"""
import json
import os
import re
import sys
import unicodedata

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, 'assets', 'crie_calcule.json')
ABA_IMPACTOS = 'Impactos do crie e calcule'
ABA_TERMICA = 'Prop. térmicas crie e calcule'

# Colunas B..L da aba de impactos
CHAVES_IMPACTO = ['gwp', 'adpf', 'adpe', 'ap', 'odp', 'epfw', 'epm', 'ept', 'pocp', 'wdp', 'ced']

# Nomes que diferem entre as duas abas (impactos -> propriedades térmicas)
ALIAS_TERMICO = {
    'perfil de aco galvanizado 9cm': 'perfil de aco galvanizado 9cm + camara de ar',
    'camara de ar 5cm': 'camara de ar',
}


def limpa(s):
    return ' '.join(str(s).split())


def chave(s):
    s = unicodedata.normalize('NFD', limpa(s)).encode('ascii', 'ignore').decode()
    return s.lower()


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def ler_termicos(ws):
    linha_het = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Heterogêneo')
    itens = {}

    for r in range(3, linha_het):
        nome = ws.cell(r, 1).value
        if nome is None or str(nome).startswith('Componente'):
            continue
        e, lam, rho, cal = (ws.cell(r, k).value for k in (3, 4, 5, 6))
        nome = limpa(nome)
        item = {'grupo': 'homogeneo', 'incluir': True, 'calculavel': True, 'motivo': None, 'camadas': [], 'avisos': [], 'notas': []}
        if not isinstance(e, (int, float)):
            item['incluir'] = False
            item['avisos'].append(f'Planilha: espessura "{e}" — camada não entra no cálculo térmico.')
        else:
            camara = chave(nome) == 'camara de ar'
            item['camadas'] = [{'material': nome, 'e_m': float(e), 'lambda': num(lam), 'rho': num(rho), 'c': num(cal),
                                'camaraAr': camara}]
            if '+ camara de ar' in chave(nome):
                item['calculavel'] = False
                item['motivo'] = ('Seção composta (perfil de aço + câmara de ar): a planilha não informa ρ, c '
                                  'nem as frações de área do perfil, necessárias ao cálculo por caminhos paralelos.')
        itens[chave(nome)] = item

    atual = None
    for r in range(linha_het + 2, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and a.startswith('http'):
            break
        if a is None and b is None:
            continue
        if a is not None:
            atual = {'grupo': 'heterogeneo', 'incluir': True, 'calculavel': True, 'motivo': None, 'camadas': [], 'avisos': [], 'notas': []}
            itens[chave(a)] = atual
        c, d, e, f, g = (ws.cell(r, k).value for k in range(3, 8))
        # Linhas dos blocos maciços estão deslocadas uma coluna: C=altura, D=espessura (m), E=λ, F=ρ, G=c
        deslocado = isinstance(e, (int, float)) and isinstance(f, (int, float)) and e < 10 and f >= 100
        if deslocado:
            e_cm, lam, rho, cal = d * 100, e, f, g
            atual['notas'].append(f'Planilha (linha {r}): colunas deslocadas; espessura lida da coluna D ({d} m).')
        else:
            e_cm, lam, rho, cal = c, d, num(e), num(f)
        atual['camadas'].append({'material': limpa(b), 'e_m': round(float(e_cm) / 100, 6), 'lambda': num(lam),
                                 'rho': num(rho), 'c': num(cal), 'camaraAr': False})
    return itens


def espessura_cm(nome, termico):
    """Espessura física exibida ao usuário."""
    nominal = re.search(r'(\d+(?:[.,]\d+)?)\s*x\s*\d+\s*x\s*\d+', nome)
    if nominal:
        return float(nominal.group(1).replace(',', '.'))
    if termico['grupo'] == 'heterogeneo':
        # A argamassa de assentamento entra em série no cálculo da planilha, mas não aumenta a largura do bloco.
        e = sum(l['e_m'] for l in termico['camadas'] if chave(l['material']) != 'argamassa')
        return round(e * 100, 2)
    if termico['camadas']:
        return round(termico['camadas'][0]['e_m'] * 100, 2)
    return 0


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser('~/Downloads/app planilha_05.09_Resumo Thiago.xlsx')
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    termicos = ler_termicos(wb[ABA_TERMICA])

    ws = wb[ABA_IMPACTOS]
    componentes = []
    for r in range(3, ws.max_row + 1):
        nome = ws.cell(r, 1).value
        if not nome:
            continue
        nome = limpa(nome)
        valores = [ws.cell(r, k).value for k in range(2, 13)]
        avisos = []
        if all(v is None for v in valores):
            avisos.append('Planilha: sem dados de impacto — considerado zero.')
        impactos = {k: (num(v) or 0.0) for k, v in zip(CHAVES_IMPACTO, valores)}

        k = chave(nome)
        termico = termicos.get(ALIAS_TERMICO.get(k, k))
        if termico is None:
            raise SystemExit(f'Sem propriedades térmicas para "{nome}"')
        avisos += termico.pop('avisos')
        notas = termico.pop('notas')

        componentes.append({
            'id': re.sub(r'[^a-z0-9]+', '-', k).strip('-'),
            'nome': nome,
            'espessura_cm': espessura_cm(nome, termico),
            'impactos': impactos,
            'termico': termico,
            'avisos': avisos,
            'notas_planilha': notas,
        })

    dados = {
        'fonte': {'arquivo': os.path.basename(xlsx), 'abas': [ABA_IMPACTOS, ABA_TERMICA]},
        'unidade_funcional': '1 m² de parede',
        'componentes': componentes,
    }
    with open(SAIDA, 'w', encoding='utf-8') as fh:
        json.dump(dados, fh, ensure_ascii=False, indent=2)
    print(f'{len(componentes)} componentes -> {SAIDA}')


if __name__ == '__main__':
    main()
